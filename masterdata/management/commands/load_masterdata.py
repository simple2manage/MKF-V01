from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from django.conf import settings
import os

from masterdata.models import *
from crops.models import *
from django.core.files import File
#
#
# class Command(BaseCommand):
#     help = 'Load initial data for Input and MachineType models from Excel file'
#
#     def handle(self, *args, **kwargs):
#         path = os.path.join(settings.BASE_DIR, 'MasterData.xlsx')
#
#         try:
#             wb = load_workbook(filename=path)
#         except Exception as e:
#             self.stdout.write(self.style.ERROR(f"Error opening Excel file: {e}"))
#             return
#
#         # ---------------- Load Inputs ---------------- #
#         if "Inputs" in wb.sheetnames:
#             sheet = wb["Inputs"]
#             for row in sheet.iter_rows(min_row=2, values_only=True):
#                 name, *rest = row
#                 if name:
#                     obj, created = Input.objects.get_or_create(name=name.strip())
#                     if created:
#                         self.stdout.write(self.style.SUCCESS(f"Input created: {name}"))
#                     else:
#                         self.stdout.write(self.style.WARNING(f"Input already exists: {name}"))
#         else:
#             self.stdout.write(self.style.WARNING("No 'Inputs' sheet found in Excel."))
#
#         # ---------------- Load MachineTypes ---------------- #
#         if "MachineTypes" in wb.sheetnames:
#             sheet = wb["MachineTypes"]
#             for row in sheet.iter_rows(min_row=2, values_only=True):
#                 name, *rest = row
#                 if name:
#                     obj, created = MachineType.objects.get_or_create(name=name.strip())
#                     if created:
#                         self.stdout.write(self.style.SUCCESS(f"MachineType created: {name}"))
#                     else:
#                         self.stdout.write(self.style.WARNING(f"MachineType already exists: {name}"))
#         else:
#             self.stdout.write(self.style.WARNING("No 'MachineTypes' sheet found in Excel."))
from django.core.management.base import BaseCommand
from openpyxl import load_workbook
from django.conf import settings
import os

from masterdata.models import InputCategory, InputSubCategory, Input, MachineType

class Command(BaseCommand):
    help = 'Load initial data for Input, InputCategory, InputSubCategory, and MachineType from Excel file'

    def handle(self, *args, **kwargs):
        path = os.path.join(settings.BASE_DIR, 'MasterData.xlsx')

        try:
            wb = load_workbook(filename=path)
        except Exception as e:
            self.stdout.write(self.style.ERROR(f"Error opening Excel file: {e}"))
            return

        # ---------------- Load Inputs ---------------- #
        if "Inputs" in wb.sheetnames:
            sheet = wb["Inputs"]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                # Assuming Excel columns: Name | Category | SubCategory
                name = row[0]
                category_name = row[1] if len(row) > 1 else None
                subcategory_name = row[2] if len(row) > 2 else None

                if not name:
                    continue

                # Get or create category
                if category_name:
                    category, _ = InputCategory.objects.get_or_create(name=category_name.strip())
                else:
                    # Default category if none provided
                    category, _ = InputCategory.objects.get_or_create(name="Default Category")

                # Get or create subcategory
                subcategory = None
                if subcategory_name:
                    subcategory, _ = InputSubCategory.objects.get_or_create(
                        category=category,
                        name=subcategory_name.strip()
                    )

                # Create Input
                obj, created = Input.objects.get_or_create(
                    name=name.strip(),
                    defaults={
                        "category": category,
                        "subcategory": subcategory,
                        "form_type": "powder",  # default
                        "unit": "kg",           # default
                    }
                )
                if created:
                    self.stdout.write(self.style.SUCCESS(f"Input created: {name}"))
                else:
                    self.stdout.write(self.style.WARNING(f"Input already exists: {name}"))

        else:
            self.stdout.write(self.style.WARNING("No 'Inputs' sheet found in Excel."))

        # ---------------- Load MachineTypes ---------------- #
        if "MachineTypes" in wb.sheetnames:
            sheet = wb["MachineTypes"]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                name = row[0]
                if not name:
                    continue

                obj, created = MachineType.objects.get_or_create(name=name.strip())
                if created:
                    self.stdout.write(self.style.SUCCESS(f"MachineType created: {name}"))
                else:
                    self.stdout.write(self.style.WARNING(f"MachineType already exists: {name}"))
        else:
            self.stdout.write(self.style.WARNING("No 'MachineTypes' sheet found in Excel."))

        #     # ---------------- Load NudgeStageImages ---------------- #
        # if "NudgeStageImages" in wb.sheetnames:
        #     sheet = wb["NudgeStageImages"]
        #     for row in sheet.iter_rows(min_row=2, values_only=True):
        #         stage, icon_filename, description = row[:3]
        #
        #         if stage:
        #             obj, created = NudgeStageImage.objects.get_or_create(stage=stage.strip())
        #
        #             # handle description
        #             if description:
        #                 obj.description = description.strip()
        #
        #             # handle icon (expects file in MEDIA_ROOT/nudge_stage_icons/)
        #             if icon_filename:
        #                 icon_path = os.path.join(settings.MEDIA_ROOT, "nudge_stage_icons", icon_filename)
        #                 if os.path.exists(icon_path):
        #                     with open(icon_path, "rb") as f:
        #                         obj.icon.save(icon_filename, File(f), save=False)
        #                 else:
        #                     self.stdout.write(self.style.WARNING(f"Icon file not found: {icon_filename}"))
        #
        #             obj.save()
        #
        #             if created:
        #                 self.stdout.write(self.style.SUCCESS(f"NudgeStageImage created: {stage}"))
        #             else:
        #                 self.stdout.write(self.style.WARNING(f"NudgeStageImage updated: {stage}"))
        # else:
        #     self.stdout.write(self.style.WARNING("No 'NudgeStageImages' sheet found in Excel."))
        # ---------------- Load NudgeStageImages ---------------- #
        if "NudgeStageImages" in wb.sheetnames:
            sheet = wb["NudgeStageImages"]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                stage, icon_path, description = row[:3]

                if stage:
                    obj, created = NudgeStageImage.objects.get_or_create(stage=stage.strip())

                    # handle description
                    if description:
                        obj.description = description.strip()

                    # handle icon (Excel now contains either full or relative path)
                    if icon_path:
                        # if relative path, make it relative to BASE_DIR
                        if not os.path.isabs(icon_path):
                            icon_path = os.path.join(settings.BASE_DIR, icon_path)

                        if os.path.exists(icon_path):
                            filename = os.path.basename(icon_path)
                            with open(icon_path, "rb") as f:
                                obj.icon.save(filename, File(f), save=False)
                        else:
                            self.stdout.write(self.style.WARNING(f"Icon file not found: {icon_path}"))

                    obj.save()

                    if created:
                        self.stdout.write(self.style.SUCCESS(f"NudgeStageImage created: {stage}"))
                    else:
                        self.stdout.write(self.style.WARNING(f"NudgeStageImage updated: {stage}"))
        else:
            self.stdout.write(self.style.WARNING("No 'NudgeStageImages' sheet found in Excel."))
